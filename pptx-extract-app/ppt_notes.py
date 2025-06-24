
import os
import time
import base64
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
from io import BytesIO
import google.generativeai as genai
from google.generativeai.types import GenerationConfig
import pytesseract
from dotenv import load_dotenv



# Load environment variables
load_dotenv()
# Configure API key
genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))
# Choose a Gemini model.
model = genai.GenerativeModel(model_name="gemini-2.0-flash-lite")


def ExtractText_LLM(image_base64: str, image_format: str = 'png', max_retries=3, delay=1, quota_refill_delay=60):
    generation_config = GenerationConfig(max_output_tokens=100)

    for attempt in range(max_retries):
        try:
            image_part = {
                'inline_data': {
                    'mime_type': f'image/{image_format}',
                    'data': image_base64
                }
            }

            result = model.generate_content(
                contents=[
                    image_part,
                    "\n",
                    "Return a very short detailed description of the image provided. "
                    "Only return one or two sentences, about 40 words. Nothing else!"
                ],
                generation_config=generation_config,
                request_options={"timeout": 10}
            )

            return result.text.strip()

        except Exception as e:
            if "Resource has been exhausted" in str(e):
                print(f"Quota exhausted, waiting {quota_refill_delay} seconds for refill...")
                time.sleep(quota_refill_delay)
            else:
                print(f"Attempt {attempt + 1} failed: {str(e)}")
                if attempt < max_retries - 1:
                    time.sleep(delay)
                else:
                    raise

def ExtractText_OCR(image_base64):
    try:
        # Decode base64 image
        image_data = base64.b64decode(image_base64)
        img = Image.open(BytesIO(image_data))

        # Extract text using OCR
        text = pytesseract.image_to_string(img)
        return text.strip()

    except Exception as e:
        print(f"Error during OCR extraction: {e}")
        return ""


def extract_images_and_add_notes(pptx_path, output_path):
    prs = Presentation(pptx_path)

    for slide_index, slide in enumerate(prs.slides):
        notes_texts = []
        image_counter = 1  # Only counts images, not all shapes

        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                image_bytes = image.blob

                # Load image using PIL
                img = Image.open(BytesIO(image_bytes))

                # Convert to base64
                buffered = BytesIO()
                image_format = img.format or "PNG"
                img.save(buffered, format=image_format)
                img_base64 = base64.b64encode(buffered.getvalue()).decode('utf-8')

                # Normalize format
                fmt = image_format.lower()
                if fmt == "jpg":
                    fmt = "jpeg"


                # Call LLM and OCR functions
                llm_note_text = ExtractText_LLM(img_base64, image_format=fmt)
                ocr_note_text = ExtractText_OCR(img_base64)

                # Append note with correct image number
                note_text = (
                    f"Image {image_counter}:\n"
                    f"LLM Text: {llm_note_text}\n"
                    f"OCR Text: {ocr_note_text}\n"
                )
                notes_texts.append(note_text)
                image_counter += 1  # Increment only for image shapes

        # Add or update notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = "\n".join(notes_texts) if notes_texts else "No images found on this slide."

    # Check if output file exists and overwrite
    if os.path.exists(output_path):
        os.remove(output_path)

    prs.save(output_path)