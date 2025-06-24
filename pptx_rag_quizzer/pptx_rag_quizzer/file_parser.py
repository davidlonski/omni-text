import uuid
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def parse_powerpoint(file_object):
    """
    Parses an in-memory PowerPoint file to extract text and images in order.

    The function extracts text from speaker notes and shapes, and image data
    from picture shapes, maintaining the slide order.

    Args:
        file_object (io.BytesIO): An in-memory byte stream of the .pptx file.

    Returns:
        list[dict]: A list of dictionaries, where each dictionary represents
                    an extracted element (text or image) with its metadata.
    """
    extracted_content = []
    prs = Presentation(file_object)

    for slide_idx, slide in enumerate(prs.slides):
        # Extract from speaker notes first
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame and slide.notes_slide.notes_text_frame.text:
            extracted_content.append({
                "id": str(uuid.uuid4()),
                "type": "text",
                "content": slide.notes_slide.notes_text_frame.text,
                "slide_number": slide_idx + 1,
                "source": "speaker_notes"
            })

        # Sort shapes by position (top-to-bottom, left-to-right) for reading order
        shapes = sorted(slide.shapes, key=lambda x: (x.top, x.left) if hasattr(x, 'top') and hasattr(x, 'left') else (0, 0))

        # Extract from shapes on the slide
        for shape in shapes:
            if shape.has_text_frame and shape.text_frame.text:
                extracted_content.append({
                    "id": str(uuid.uuid4()),
                    "type": "text",
                    "content": shape.text_frame.text,
                    "slide_number": slide_idx + 1,
                    "source": "slide_shape"
                })
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_bytes = shape.image.blob
                image_ext = shape.image.ext
                extracted_content.append({
                    "id": str(uuid.uuid4()),
                    "type": "image",
                    "content": image_bytes,
                    "extension": image_ext,
                    "slide_number": slide_idx + 1,
                    "source": "slide_image"
                })
    return extracted_content
