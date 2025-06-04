import os
import zipfile
from pathlib import Path
from models import Image, Text, ImageStatus
from docx import Document as DocxDocument
from pptx import Presentation
import xml.etree.ElementTree as ET


############################################
# Enhanced extraction with position tracking #
############################################
def extract_document_structure(file_path, output_dir="extracted_images"):
    """
    Extract images and text from office documents with position tracking.
    Returns structured data for document reconstruction.
    """
    ext = Path(file_path).suffix.lower()
    os.makedirs(output_dir, exist_ok=True)
    
    # Initialize return structures
    images_list = []
    text_list = []
    document_structure = []  # Ordered list of elements with positions
    
    if ext == ".docx":
        return extract_docx_structure(file_path, output_dir)
    elif ext == ".pptx":
        return extract_pptx_structure(file_path, output_dir)
    else:
        raise ValueError("Unsupported file type. Only .pptx and .docx are supported.")

def extract_docx_structure(file_path, output_dir="extracted_images"):
    """Extract structured content from DOCX files"""
    images_list = []
    text_list = []
    document_structure = []
    
    # Extract images first
    extracted_image_paths = extract_images_from_zip(file_path, "word/media/", output_dir)
    
    # Parse document for text and element positions
    doc = DocxDocument(file_path)
    
    element_order = 0
    current_page = 1  # DOCX doesn't have clear page breaks, so we estimate
    
    for paragraph_idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip():  # Only process non-empty paragraphs
            text_obj = Text(
                text=paragraph.text,
                page_number=current_page,
                status=ImageStatus.PENDING
            )
            text_list.append(text_obj)
            
            document_structure.append({
                'type': 'text',
                'element_order': element_order,
                'page_number': current_page,
                'paragraph_index': paragraph_idx,
                'object_id': text_obj.id,
                'content_preview': paragraph.text[:50] + "..." if len(paragraph.text) > 50 else paragraph.text
            })
            element_order += 1
        
        # Check for page breaks (rough estimation)
        if len(paragraph.text) > 1000:
            current_page += 1
    
    # Process extracted images and associate with document structure
    for img_idx, img_path in enumerate(extracted_image_paths):
        # Estimate page placement (could be improved with more sophisticated parsing)
        estimated_page = (img_idx // 2) + 1  # Rough estimation
        
        image_obj = Image(
            page_number=estimated_page,
            file_path=img_path,
            status=ImageStatus.PENDING
        )
        images_list.append(image_obj)
        
        document_structure.append({
            'type': 'image',
            'element_order': element_order,
            'page_number': estimated_page,
            'image_index': img_idx,
            'object_id': image_obj.id,
            'file_path': img_path,
            'content_preview': f"Image {img_idx + 1}"
        })
        element_order += 1
    
    # Sort document structure by element order
    document_structure.sort(key=lambda x: x['element_order'])
    
    return images_list, text_list, document_structure

def extract_pptx_structure(file_path, output_dir="extracted_images"):
    """Extract structured content from PPTX files"""
    images_list = []
    text_list = []
    document_structure = []
    
    # Extract images first
    extracted_image_paths = extract_images_from_zip(file_path, "ppt/media/", output_dir)
    
    # Parse presentation for text and slide structure
    presentation = Presentation(file_path)
    
    element_order = 0
    images_per_slide = {}
    
    # First pass: extract text from each slide
    for slide_idx, slide in enumerate(presentation.slides):
        slide_number = slide_idx + 1
        slide_text_elements = []
        
        for shape_idx, shape in enumerate(slide.shapes):
            if hasattr(shape, "text") and shape.text.strip():
                text_obj = Text(
                    text=shape.text,
                    page_number=slide_number,
                    status=ImageStatus.PENDING
                )
                text_list.append(text_obj)
                slide_text_elements.append(text_obj)
                
                document_structure.append({
                    'type': 'text',
                    'element_order': element_order,
                    'page_number': slide_number,
                    'slide_index': slide_idx,
                    'shape_index': shape_idx,
                    'object_id': text_obj.id,
                    'content_preview': shape.text[:50] + "..." if len(shape.text) > 50 else shape.text
                })
                element_order += 1
        
        # Estimate images per slide (distribute evenly)
        images_per_slide[slide_number] = len(extracted_image_paths) // len(presentation.slides)
    
    # Second pass: associate images with slides
    img_counter = 0
    for slide_idx, slide in enumerate(presentation.slides):
        slide_number = slide_idx + 1
        estimated_images_on_slide = images_per_slide.get(slide_number, 0)
        
        # Add images estimated to be on this slide
        for i in range(estimated_images_on_slide):
            if img_counter < len(extracted_image_paths):
                img_path = extracted_image_paths[img_counter]
                
                image_obj = Image(
                    page_number=slide_number,
                    file_path=img_path,
                    status=ImageStatus.PENDING
                )
                images_list.append(image_obj)
                
                document_structure.append({
                    'type': 'image',
                    'element_order': element_order,
                    'page_number': slide_number,
                    'slide_index': slide_idx,
                    'image_index': img_counter,
                    'object_id': image_obj.id,
                    'file_path': img_path,
                    'content_preview': f"Image on Slide {slide_number}"
                })
                element_order += 1
                img_counter += 1
    
    # Add any remaining images to the last slide
    while img_counter < len(extracted_image_paths):
        img_path = extracted_image_paths[img_counter]
        last_slide = len(presentation.slides)
        
        image_obj = Image(
            page_number=last_slide,
            file_path=img_path,
            status=ImageStatus.PENDING
        )
        images_list.append(image_obj)
        
        document_structure.append({
            'type': 'image',
            'element_order': element_order,
            'page_number': last_slide,
            'slide_index': last_slide - 1,
            'image_index': img_counter,
            'object_id': image_obj.id,
            'file_path': img_path,
            'content_preview': f"Additional Image {img_counter + 1}"
        })
        element_order += 1
        img_counter += 1
    
    # Sort document structure by page number and element order
    document_structure.sort(key=lambda x: (x['page_number'], x['element_order']))
    
    return images_list, text_list, document_structure

def extract_images_from_zip(zip_path, media_folder, output_dir):
    """Helper function to extract images from zip-based office files"""
    extracted_image_paths = []
    
    with zipfile.ZipFile(zip_path, 'r') as archive:
        for file in archive.namelist():
            if file.startswith(media_folder) and file.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.tif', '.svg')):
                image_data = archive.read(file)
                image_name = os.path.basename(file)
                rel_path = os.path.join(output_dir, image_name)
                with open(rel_path, 'wb') as img_file:
                    img_file.write(image_data)
                extracted_image_paths.append(rel_path.replace("\\", "/"))
    
    return extracted_image_paths

def get_document_layout_summary(document_structure):
    """Generate a summary of document layout for reconstruction"""
    layout_summary = {
        'total_elements': len(document_structure),
        'pages': {},
        'element_types': {'text': 0, 'image': 0}
    }
    
    for element in document_structure:
        page_num = element['page_number']
        element_type = element['type']
        
        # Count by page
        if page_num not in layout_summary['pages']:
            layout_summary['pages'][page_num] = {'text': 0, 'image': 0, 'elements': []}
        
        layout_summary['pages'][page_num][element_type] += 1
        layout_summary['pages'][page_num]['elements'].append({
            'type': element_type,
            'order': element['element_order'],
            'preview': element['content_preview']
        })
        
        # Count by type
        layout_summary['element_types'][element_type] += 1
    
    return layout_summary

############################################
# Original function (kept for compatibility) #
############################################
def extract_images_from_office(file_path, output_dir="extracted_images"):
    ext = Path(file_path).suffix.lower()
    os.makedirs(output_dir, exist_ok=True)
    extracted_image_paths = []

    def extract_from_zip(zip_path, media_folder):
        with zipfile.ZipFile(zip_path, 'r') as archive:
            for file in archive.namelist():
                if file.startswith(media_folder) and file.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.tif', '.bmp', '.svg')):
                    image_data = archive.read(file)
                    image_name = os.path.basename(file)
                    rel_path = os.path.join(output_dir, image_name)
                    with open(rel_path, 'wb') as img_file:
                        img_file.write(image_data)
                    extracted_image_paths.append(rel_path.replace("\\", "/"))

    if ext == ".pptx":
        extract_from_zip(file_path, "ppt/media/")
    elif ext == ".docx":
        extract_from_zip(file_path, "word/media/")
    else:
        raise ValueError("Unsupported file type. Only .pptx and .docx are supported.")

    return extracted_image_paths

# Example usage function
def test_document_extraction(file_path):
    """Test the new extraction function"""
    try:
        images_list, text_list, document_structure = extract_document_structure(file_path)
        layout_summary = get_document_layout_summary(document_structure)
        
        print(f"Extracted {len(images_list)} images and {len(text_list)} text elements")
        print(f"Document has {len(layout_summary['pages'])} pages")
        
        for page_num, page_info in layout_summary['pages'].items():
            print(f"Page {page_num}: {page_info['text']} text elements, {page_info['image']} images")
        
        return images_list, text_list, document_structure, layout_summary
        
    except Exception as e:
        print(f"Error extracting document: {e}")
        return [], [], [], {}


